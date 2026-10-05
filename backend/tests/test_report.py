"""Unit tests for deterministic DOCX review report generator (services/report.py)."""

from __future__ import annotations

import io
import re
from typing import cast
from uuid import uuid4

import docx
import fitz

from domain.enums import DocumentStatus, IssueCategory, ReportLanguage, Severity
from models.document import Document
from models.issue import Issue
from schemas.budinski import MajorFinding
from services.export import assessment_from_document_findings, generate_ale_review_docx
from services.report import build_review_report
from services.report_locale import humanize_finding_type
from services.report_synthesizer import ReportSynthesizer


def _docx_text(word_doc: docx.document.Document) -> str:
    return "\n".join(
        [p.text for p in word_doc.paragraphs]
        + [cell.text for table in word_doc.tables for row in table.rows for cell in row.cells]
    )


def test_build_review_report_empty_findings() -> None:
    """DOCX report generates cleanly when there are zero included findings."""
    doc = Document(
        id=uuid4(),
        original_filename="CLEAN_DOC.pdf",
        safe_filename="clean_doc",
        status=DocumentStatus.COMPLETED,
    )
    docx_bytes = build_review_report(doc, [])
    assert len(docx_bytes) > 0

    word_doc = docx.Document(io.BytesIO(docx_bytes))
    full_text = "\n".join(p.text for p in word_doc.paragraphs)
    assert "DOCUMENT REVIEW ENGINEERING" in full_text
    assert "CLEAN_DOC.pdf" in full_text
    assert "0 findings" in full_text
    assert "Blockers" in full_text
    assert "Next revision findings" in full_text


def test_build_review_report_full_composition() -> None:
    """DOCX report categorizes blockers, other findings, and language findings."""
    doc = uuid4()
    document = Document(
        id=doc,
        original_filename="VESSEL_CALC.pdf",
        safe_filename="vessel_calc",
        status=DocumentStatus.COMPLETED,
    )

    issue_blocker = Issue(
        id=uuid4(),
        document_id=doc,
        category=IssueCategory.TRACEABILITY,
        type="TABLE_MATH_MISMATCH",
        severity=Severity.CRITICAL,
        message="Total 1500 != computed 1600",
        page_number=3,
        evidence={"suggestion": "Adjust row 4 to match sum"},
        included_in_report=True,
        reviewer_note="Confirmed with senior pressure vessel engineer.",
    )

    issue_other = Issue(
        id=uuid4(),
        document_id=doc,
        category=IssueCategory.TRACEABILITY,
        type="REFERENCE_DRIFT",
        severity=Severity.MEDIUM,
        message="Section 3.2 references Appendix B which was moved to Appendix C",
        page_number=12,
        evidence={},
        included_in_report=True,
        reviewer_note=None,
    )

    issue_lang = Issue(
        id=uuid4(),
        document_id=doc,
        category=IssueCategory.LINGUISTIC,
        type="SPELLING_ERROR",
        severity=Severity.LOW,
        message="Typo 'tempature'",
        page_number=5,
        evidence={"suggestion": "temperature"},
        included_in_report=True,
        reviewer_note="Reviewer approved typo fix.",
    )

    issue_excluded = Issue(
        id=uuid4(),
        document_id=doc,
        category=IssueCategory.LINGUISTIC,
        type="AMBIGUITY",
        severity=Severity.INFO,
        message="Ambiguous phrasing",
        page_number=8,
        evidence={},
        included_in_report=False,
        reviewer_note="Not relevant for this rev",
    )

    docx_bytes = build_review_report(
        document,
        [issue_blocker, issue_other, issue_lang, issue_excluded],
    )
    assert len(docx_bytes) > 0

    word_doc = docx.Document(io.BytesIO(docx_bytes))
    full_text = "\n".join(p.text for p in word_doc.paragraphs)

    assert "VESSEL_CALC.pdf" in full_text
    all_table_text = "\n".join(
        " ".join(cell.text for cell in row.cells) for table in word_doc.tables for row in table.rows
    )
    assert "Arithmetic Summation Discrepancy" in _docx_text(word_doc)
    assert "Confirmed with senior pressure vessel engineer." in all_table_text
    assert "Adjust row 4 to match sum" in all_table_text
    assert "Cross-Reference Pagination Drift" in all_table_text
    assert "SPELLING_ERROR" not in all_table_text
    assert "temperature" in all_table_text

    # Excluded finding must NOT be present
    assert "Ambiguous phrasing" not in full_text

    # Verify scorecard table
    table = word_doc.tables[0]
    data = {row.cells[0].text: row.cells[1].text for row in table.rows[1:]}
    assert data["Included findings"] == "3"
    assert data["Blockers"] == "1"
    assert data["Language and mechanics"] == "1"
    assert data["Other consistency findings"] == "1"


def test_finding_identifiers_are_humanized_in_executive_exports() -> None:
    from services.report_locale import humanize_finding_type

    assert humanize_finding_type("TABLE_MATH_MISMATCH") == "Arithmetic Summation Discrepancy"
    assert (
        humanize_finding_type("TABLE_MATH_MISMATCH", count=2)
        == "Arithmetic Summation Discrepancies"
    )
    assert humanize_finding_type("RULE_GRAMMAR_SPLIT_WORD") == "Broken Word / Erroneous Hyphenation"
    assert (
        humanize_finding_type("MISSING_DOCUMENT_NUMBER", language=ReportLanguage.INDONESIAN)
        == "Metadata Kendali Dokumen Tidak Lengkap"
    )

    document_id = uuid4()
    issue = Issue(
        id=uuid4(),
        document_id=document_id,
        category=IssueCategory.TRACEABILITY,
        type="TABLE_MATH_MISMATCH",
        severity=Severity.CRITICAL,
        message="TABLE_MATH_MISMATCH: Total 1500 does not match the row sum 1600.",
        page_number=3,
        evidence={"label": "Table 4.2", "suggestion": "Reconcile the total with the listed rows."},
        included_in_report=True,
    )
    report = docx.Document(
        io.BytesIO(
            generate_ale_review_docx(
                assessment_from_document_findings(
                    Document(
                        id=document_id,
                        original_filename="sample.pdf",
                        safe_filename="sample",
                        status=DocumentStatus.COMPLETED,
                    ),
                    [issue],
                )
            )
        )
    )
    text = _docx_text(report)
    assert "TABLE_MATH_MISMATCH" not in text
    assert "Finding code" not in text
    assert "Arithmetic Summation Discrepancy" in text
    assert "Table 4.2" in text

    group = ReportSynthesizer().synthesize_blocker_groups([issue])[0]
    assert "TABLE_MATH_MISMATCH" not in group["message"]
    assert "Arithmetic Summation Discrepancy" in group["message"]


def test_grouped_narrative_and_low_severity_language_findings_are_explicit() -> None:
    document_id = uuid4()
    findings = [
        Issue(
            id=uuid4(),
            document_id=document_id,
            category=IssueCategory.TRACEABILITY,
            type="MISSING_TARGET",
            severity=Severity.CRITICAL,
            message="Unresolved target at page 1.",
            page_number=1,
            evidence={},
            included_in_report=True,
        ),
        Issue(
            id=uuid4(),
            document_id=document_id,
            category=IssueCategory.TRACEABILITY,
            type="REF_DRIFT",
            severity=Severity.CRITICAL,
            message="Reference pagination differs.",
            page_number=3,
            evidence={},
            included_in_report=True,
        ),
        *[
            Issue(
                id=uuid4(),
                document_id=document_id,
                category=IssueCategory.TRACEABILITY,
                type="TABLE_MATH_MISMATCH",
                severity=Severity.CRITICAL,
                message=f"Table total {idx} does not reconcile.",
                page_number=page,
                evidence={},
                included_in_report=True,
            )
            for idx, page in enumerate((4, 4, 11), start=1)
        ],
        Issue(
            id=uuid4(),
            document_id=document_id,
            category=IssueCategory.LINGUISTIC,
            type="RULE_GRAMMAR_SPLIT_WORD",
            severity=Severity.LOW,
            message="Broken word detected.",
            page_number=22,
            evidence={"original_text": "nine-teen", "suggestion": "nineteen"},
            included_in_report=True,
        ),
        Issue(
            id=uuid4(),
            document_id=document_id,
            category=IssueCategory.LINGUISTIC,
            type="RULE_GRAMMAR_UNCLOSED_SENTENCE",
            severity=Severity.LOW,
            message="Sentence is incomplete.",
            page_number=29,
            evidence={
                "original_text": "classified as Priority by end-2027",
                "suggestion": "Complete the sentence.",
            },
            included_in_report=True,
        ),
    ]
    document = Document(
        id=document_id,
        original_filename="reviewed.pdf",
        safe_filename="reviewed",
        status=DocumentStatus.COMPLETED,
    )
    assessment = assessment_from_document_findings(document, findings)
    blocker_text = " ".join(item.what_it_says for item in assessment.blockers)
    assert (
        "The assessment identified 5 document-control findings encompassing "
        "1 Unresolved Reference Target, 1 Cross-Reference Pagination Drift, and "
        "3 Arithmetic Summation Discrepancies across printed pages 1, 3, 4, and 11."
    ) in blocker_text
    assert len(assessment.language_findings) == 2

    export_text = _docx_text(docx.Document(io.BytesIO(generate_ale_review_docx(assessment))))
    assert "Page 22: Broken Word / Erroneous Hyphenation" in export_text
    assert "‘nine-teen’" in export_text
    assert "Page 29: Incomplete / Truncated Sentence" in export_text
    assert "‘classified as Priority by end-2027’" in export_text
    assert "No language or mechanics findings identified." not in export_text
    assert "MISSING_TARGET" not in export_text
    assert "RULE_GRAMMAR_UNCLOSED_SENTENCE" not in export_text

    concise_assessment = assessment_from_document_findings(
        document,
        findings,
        include_language_mechanics=False,
    )
    concise_text = _docx_text(
        docx.Document(
            io.BytesIO(
                generate_ale_review_docx(
                    concise_assessment,
                    include_language_mechanics=False,
                )
            )
        )
    )
    assert concise_assessment.language_findings == []
    assert "Language and mechanics, by page" not in concise_text
    assert "nine-teen" not in concise_text


def test_repeated_language_findings_are_grouped_across_pages() -> None:
    synthesizer = ReportSynthesizer()
    findings = [
        {
            "page": "2",
            "finding_type": "SPELLING_ERROR",
            "as_written": "teh",
            "suggested": "the",
        },
        {
            "page": "8",
            "finding_type": "SPELLING_ERROR",
            "as_written": "teh",
            "suggested": "the",
        },
        {
            "page": "9",
            "finding_type": "SPELLING_ERROR",
            "as_written": "teh",
            "suggested": "the",
        },
    ]

    rows = synthesizer.aggregate_language_findings(findings)
    narrative = synthesizer.synthesize_language_narrative(findings)

    assert len(rows) == 1
    assert rows[0]["page"] == "2, 8, 9"
    assert "3 occurrences" in rows[0]["finding_type"]
    assert len(narrative) == 1
    assert "Pages 2, 8, 9" in narrative[0]
    assert synthesizer.aggregate_language_findings(
        findings,
        include_language_mechanics=False,
    ) == []
    assert synthesizer.synthesize_language_mechanics(
        findings,
        include_language_mechanics=False,
    ) == []


def test_bottom_line_and_visual_finding_table_humanize_raw_categories() -> None:
    synthesizer = ReportSynthesizer()
    assert synthesizer.generate_bottom_line([{"title": "TRACEABILITY (6 instances)"}]) == (
        "Resolve 6 Critical Traceability and Control Discrepancies before document reissue."
    )

    document_id = uuid4()
    issue = Issue(
        id=uuid4(),
        document_id=document_id,
        category=IssueCategory.TRACEABILITY,
        type="TRACEABILITY",
        severity=Severity.CRITICAL,
        message="Document-control traceability issue.",
        page_number=2,
        evidence={},
        included_in_report=True,
    )
    document = Document(
        id=document_id,
        original_filename="review.pdf",
        safe_filename="review",
        status=DocumentStatus.COMPLETED,
    )
    assessment = assessment_from_document_findings(document, [issue])
    assessment.blockers[0].title = "TRACEABILITY (6 instances)"
    assessment.bottom_line = "Resolve TRACEABILITY (6 instances)"
    assessment.major_findings = [
        MajorFinding(
            number=1,
            finding="TABLE_MATH_MISMATCH: Total does not reconcile.",
            what_would_fix_it="Correct TABLE_MATH_MISMATCH in the summary.",
        )
    ]
    report = docx.Document(io.BytesIO(generate_ale_review_docx(assessment)))
    text = _docx_text(report)
    assert (
        "Resolve 6 Critical Traceability and Control Discrepancies before document reissue." in text
    )
    assert "TRACEABILITY" not in text
    assert "TABLE_MATH_MISMATCH" not in text

    finding_table = next(
        table
        for table in report.tables
        if any(cell.text == "Finding" for cell in table.rows[0].cells)
    )
    visual_rows = [cell.text for row in finding_table.rows[1:] for cell in row.cells]
    assert all("_" not in cell for cell in visual_rows)
    assert not re.search(r"\b[A-Z]{2,}(?:_[A-Z0-9]+)*\b", " ".join(visual_rows))
    assert "Arithmetic Summation Discrepancy: Total does not reconcile." in " ".join(visual_rows)


def test_report_types_are_humanized_in_docx_and_mock_compiled_pdf() -> None:
    assert humanize_finding_type("table_math_mismatch", count=3) == (
        "Arithmetic Summation Discrepancies"
    )
    assert humanize_finding_type("spelling_split", count=1) == (
        "Broken Word / Erroneous Hyphenation"
    )

    document_id = uuid4()
    document = Document(
        id=document_id,
        original_filename="finding-label-check.pdf",
        safe_filename="finding-label-check",
        status=DocumentStatus.COMPLETED,
    )
    issues = [
        Issue(
            id=uuid4(),
            document_id=document_id,
            category=IssueCategory.TRACEABILITY,
            type="REF_DRIFT",
            severity=Severity.CRITICAL,
            message="The cross-reference page differs from the target page.",
            page_number=3,
            evidence={},
            included_in_report=True,
        ),
        Issue(
            id=uuid4(),
            document_id=document_id,
            category=IssueCategory.LAYOUT,
            type="cross_page_break",
            severity=Severity.MEDIUM,
            message="The table is split across pages.",
            page_number=4,
            evidence={},
            included_in_report=True,
        ),
        Issue(
            id=uuid4(),
            document_id=document_id,
            category=IssueCategory.BUDINSKI,
            type="budinski_rework",
            severity=Severity.MEDIUM,
            message="The conclusion needs synthesis.",
            page_number=8,
            evidence={},
            included_in_report=True,
        ),
    ]
    report_bytes = generate_ale_review_docx(assessment_from_document_findings(document, issues))
    report = docx.Document(io.BytesIO(report_bytes))
    docx_text = _docx_text(report)

    def mock_compile_pdf(text: str) -> bytes:
        pdf = fitz.open()
        chunks = [text[index : index + 1800] for index in range(0, len(text), 1800)]
        for chunk in chunks:
            page = pdf.new_page(width=612, height=792)
            page.insert_textbox(fitz.Rect(36, 36, 576, 756), chunk, fontname="helv", fontsize=7)
        return cast(bytes, pdf.tobytes())

    mock_pdf_bytes = mock_compile_pdf(docx_text)
    with fitz.open(stream=mock_pdf_bytes, filetype="pdf") as mock_pdf:
        pdf_text = "\n".join(page.get_text() for page in mock_pdf)

    forbidden_types = ("ref_drift", "REF_DRIFT", "cross_page_break", "budinski_rework")
    for raw_type in forbidden_types:
        assert raw_type not in docx_text
        assert raw_type not in pdf_text


def _reference_rule_issue(
    doc: object = None,
    standard: str = "ASME BPVC.VIII.1",
) -> Issue:
    from uuid import UUID

    return Issue(
        id=uuid4(),
        document_id=UUID(int=0),
        category=IssueCategory.TRACEABILITY,
        type="REFERENCE_RULE",
        severity=Severity.HIGH,
        message="Hydrostatic test ratio below the mandatory UG-99(b) factor.",
        page_number=4,
        evidence={
            "kind": "REFERENCE_RULE",
            "standard": standard,
            "edition": "2021",
            "clause": "UG-99(b)",
            "standard_page": 76,
            "rule_kind": "numeric_limit",
            "detected_parameter": "Hydrostatic test pressure ratio",
            "detected_value": "1.15",
            "location": {
                "page_index": 3,
                "x0": 0.0,
                "y0": 0.0,
                "x1": 612.0,
                "y1": 792.0,
                "page_width": 612.0,
                "page_height": 792.0,
            },
        },
        included_in_report=True,
        reviewer_note=None,
    )


def test_reference_findings_are_compact() -> None:
    """Reference findings stay in the compact action table."""
    document = Document(
        id=uuid4(),
        original_filename="VESSEL_SPEC.pdf",
        safe_filename="vessel_spec",
        status=DocumentStatus.COMPLETED,
    )
    governing = _reference_rule_issue(uuid4(), standard="ASME BPVC.VIII.1")
    company = _reference_rule_issue(uuid4(), standard="Company Spec CS-101 Addenda")
    company.evidence = dict(company.evidence, compliance_status="UNRESOLVED")

    docx_bytes = build_review_report(document, [governing, company])
    assert len(docx_bytes) > 0

    word_doc = docx.Document(io.BytesIO(docx_bytes))
    full_text = "\n".join(p.text for p in word_doc.paragraphs)
    assert "Blockers" in full_text
    assert "Next revision findings" in full_text

    all_table_text = "\n".join(
        " ".join(cell.text for cell in row.cells) for t in word_doc.tables for row in t.rows
    )
    assert "Governing Industry Standards" not in full_text

    # Column data and text-only standard locator
    assert "ASME BPVC.VIII.1" in all_table_text
    assert "Clause UG-99(b), standard page 76" in all_table_text
    assert "document page 4" in all_table_text
    assert "1.15" in all_table_text

    assert "ASME BPVC.VIII.1" in all_table_text
    assert "Company Spec CS-101 Addenda" in all_table_text


def test_report_rejects_invalid_bounding_boxes() -> None:
    """Evidence boxes outside the scanned page geometry are dropped, DOCX still valid."""
    document = Document(
        id=uuid4(),
        original_filename="BROKEN_BOX.pdf",
        safe_filename="broken_box",
        status=DocumentStatus.COMPLETED,
    )
    bad = _reference_rule_issue(uuid4())
    bad.evidence = dict(
        bad.evidence,
        location={
            "page_index": 0,
            "x0": -5.0,
            "y0": 0.0,
            "x1": 9999.0,
            "y1": 792.0,
            "page_width": 612.0,
            "page_height": 792.0,
        },
    )

    docx_bytes = build_review_report(document, [bad])
    assert len(docx_bytes) > 0

    word_doc = docx.Document(io.BytesIO(docx_bytes))
    all_table_text = "\n".join(
        " ".join(cell.text for cell in row.cells) for t in word_doc.tables for row in t.rows
    )
    assert "document evidence on page" not in all_table_text
    assert "Clause UG-99(b), standard page 76" in all_table_text


def test_report_orders_severity_and_uses_printed_reference_labels() -> None:
    """Compact findings are ordered by severity, not PDF page sequence."""
    document = Document(
        id=uuid4(),
        original_filename="PAGINATED_REPORT.pdf",
        safe_filename="paginated_report",
        status=DocumentStatus.COMPLETED,
    )
    low = Issue(
        id=uuid4(),
        document_id=document.id,
        category=IssueCategory.LINGUISTIC,
        type="DICTIONARY_TERM",
        severity=Severity.LOW,
        message="Term should be checked",
        page_number=1,
        evidence={
            "kind": "LINGUISTIC",
            "original_text": "term",
            "location": {"page_index": 0, "x0": 1, "y0": 1, "x1": 2, "y1": 2},
        },
        included_in_report=True,
    )
    high = Issue(
        id=uuid4(),
        document_id=document.id,
        category=IssueCategory.TRACEABILITY,
        type="REFERENCE_DRIFT",
        severity=Severity.HIGH,
        message="TOC target does not match",
        page_number=99,
        evidence={
            "kind": "REFERENCE_DRIFT",
            "label": "Section 4",
            "referenced_page_label": "21",
            "actual_page_label": "23",
            "what_it_says": "Section 4",
        },
        included_in_report=True,
    )

    word_doc = docx.Document(io.BytesIO(build_review_report(document, [low, high])))
    findings_tables = [
        table for table in word_doc.tables if table.rows[0].cells[0].text == "Severity"
    ]
    blocker_rows = [[cell.text for cell in row.cells] for row in findings_tables[0].rows[1:]]
    next_rows = [[cell.text for cell in row.cells] for row in findings_tables[1].rows[1:]]
    assert blocker_rows[0][0] == "High"
    assert "printed pages 21" in blocker_rows[0][2]
    assert next_rows[0][0] == "Low"


def test_export_uses_document_identity_without_sample_fallbacks() -> None:
    document = Document(
        id=uuid4(),
        original_filename="Inspection_Report.pdf",
        safe_filename="inspection_report",
        status=DocumentStatus.COMPLETED,
        extracted_metadata={
            "document_number": "ID-55",
            "revision": "B",
            "document_date": "2 Oct 2026",
        },
    )
    issue = Issue(
        id=uuid4(),
        document_id=document.id,
        category=IssueCategory.TRACEABILITY,
        type="SECTION_EVIDENCE",
        severity=Severity.HIGH,
        message="Objective evidence is present in Section 2.",
        page_number=2,
        evidence={"section": "Objective"},
        included_in_report=True,
    )
    assessment = assessment_from_document_findings(document, [issue])
    assert assessment.scorecard.blockers_count == 1
    assert assessment.scorecard.majors_count == 0
    assert assessment.scorecard.minors_count == 0
    assert "doc=ID-55 | rev=B | date=2 Oct 2026" in assessment.review_score_string
    assert "blockers=1 majors=0 minors=0" in assessment.review_score_string
    text = _docx_text(docx.Document(io.BytesIO(generate_ale_review_docx(assessment))))
    assert "Inspection_Report.pdf" in text
    assert "Objective evidence is present in Section 2." in text
    for prohibited in ("ALE roadmap", "DOC-001", "Rev A", "2026-08-21", "MEPG", "Grissik"):
        assert prohibited not in text


def test_structural_findings_use_ranges_and_source_excerpt() -> None:
    synthesizer = ReportSynthesizer()
    findings = [
        {
            "category": "LAYOUT",
            "type": "UNCONTROLLED_PAGE",
            "page_number": page,
            "message": "Repeated layout control finding.",
            "evidence": {"original_text": "Source excerpt from the affected page."},
            "suggestion": "Apply the controlled document template.",
        }
        for page in (2, 3, 4, 8)
    ]

    rows = synthesizer.synthesize_blocker_groups(findings)

    assert len(rows) == 1
    assert "printed pages 2, 3, 4, and 8" in rows[0]["message"]
    assert rows[0]["excerpt"] == "Source excerpt from the affected page."
    assert "Repeated layout control finding." not in rows[0]["excerpt"]


def test_review_score_string_uses_current_metadata_averages_and_counts() -> None:
    synthesizer = ReportSynthesizer()
    score = synthesizer.generate_review_score_string(
        {"document_number": "DOC-42", "revision": "C", "document_date": "2 Oct 2026"},
        {
            "group_averages": {
                "Group I": 4,
                "Group II": 3.25,
                "Group III": 2.5,
                "Group IV": 4.125,
            },
            "baseline_score": 3,
        },
        blockers=2,
        majors=5,
        minors=7,
    )

    assert score == (
        "REVIEWSCORE | doc=DOC-42 | rev=C | date=2 Oct 2026 | "
        "I=4.00 II=3.25 III=2.50 IV=4.12 | baseline=3/4 | "
        "blockers=2 majors=5 minors=7"
    )


def test_demonstration_rewrite_synthesizes_six_to_eight_source_grounded_sentences() -> None:
    result = ReportSynthesizer().generate_demonstration_rewrite(
        {
            "introduction": ["The purpose of this assessment is to evaluate vessel integrity."],
            "results": ["The inspection found a minimum wall thickness of 8.2 mm."],
            "conclusions": ["The vessel remains suitable for continued service."],
            "recommendations": [{"text": "Owner: Integrity team to reinspect by 2 Oct 2027."}],
        },
        [],
    )

    assert 6 <= len(result["demonstration_items"]) <= 8
    rewrite = " ".join(result["demonstration_items"])
    assert "8.2 mm" in rewrite
    assert "Owner: Integrity team" in rewrite
    assert "is inferred in this demonstration" in rewrite
