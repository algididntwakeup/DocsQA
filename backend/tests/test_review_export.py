"""
Unit and integration tests for Budinski DOCX Review Report Export Service.
========================================================================
Tests 'generate_ale_review_docx' and 'create_review_report_template' against
the full 12-page Review-ALE baseline specification and Kenneth G. Budinski's
Appendix 12 checklist.
"""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any

import docx
import pytest
from docx.document import Document as DocxDocument

from domain.enums import ReportLanguage
from schemas.budinski import (
    AssessmentData,
    AssessmentMetadata,
    BaselineMeasures,
    BudinskiScorecard,
    ConclusionsAndCraftGroup,
    ReportMechanicsGroup,
    ScoreItem,
    StyleGroup,
    TechnicalContentGroup,
)
from services.docx_styler import create_callout_box, format_table_header, set_cell_shading
from services.export import (
    LibreOfficeConversionError,
    assessment_from_document_findings,
    convert_docx_to_pdf,
    create_review_report_template,
    generate_ale_review_docx,
)
from services.report_synthesizer import ReportSynthesizer
from tests.canonical_ale_fixture import create_canonical_ale_assessment_data


def _extract_all_doc_text(doc: DocxDocument) -> str:
    """Extract all text from paragraphs, tables, and callouts in a Document."""
    parts: list[str] = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if p.text.strip():
                        parts.append(p.text)
    return "\n".join(parts)


def _section_position(text: str, heading: str) -> int:
    return text.index(heading)


@pytest.fixture
def canonical_assessment() -> AssessmentData:
    """Fixture providing complete canonical ALE assessment data."""
    return create_canonical_ale_assessment_data()


class TestReviewExportDocx:
    """Verification suite for generate_ale_review_docx."""

    def test_generate_ale_review_docx_returns_valid_bytes(
        self,
        canonical_assessment: AssessmentData,
    ) -> None:
        """Report generation succeeds and returns valid, non-empty DOCX binary stream."""
        docx_bytes = generate_ale_review_docx(canonical_assessment)
        assert isinstance(docx_bytes, bytes)
        assert len(docx_bytes) > 5000
        # Check ZIP / DOCX magic bytes (PK\x03\x04)
        assert docx_bytes.startswith(b"PK\x03\x04")

        # Document can be parsed by python-docx
        doc = docx.Document(io.BytesIO(docx_bytes))
        assert len(doc.paragraphs) > 0
        assert len(doc.tables) > 0

    def test_document_identity_and_metadata_block(
        self,
        canonical_assessment: AssessmentData,
    ) -> None:
        """Section 1: Document title, subtitle, and metadata block are correctly rendered."""
        docx_bytes = generate_ale_review_docx(canonical_assessment)
        doc = docx.Document(io.BytesIO(docx_bytes))
        full_text = _extract_all_doc_text(doc)

        # Title and Subtitle
        assert canonical_assessment.title in full_text
        assert canonical_assessment.subtitle in full_text

        # Metadata table fields
        assert "Document reviewed" in full_text
        assert canonical_assessment.metadata.document_reviewed in full_text
        assert "Type of review" in full_text
        assert "Technical writing review only" in full_text
        assert "Basis" in full_text
        assert "Appendix 12 checklist" in full_text
        assert "Scoring" in full_text
        assert "1 = disagree, 5 = agree" in full_text
        assert "Note" in full_text
        assert "The cover names the reviewer" in full_text
        assert "Not covered" in full_text
        assert "The engineering itself" in full_text

    def test_summary_judgement_and_bottom_line(
        self,
        canonical_assessment: AssessmentData,
    ) -> None:
        """Section 2: Summary judgement text and callout box 'BOTTOM LINE' rendered."""
        docx_bytes = generate_ale_review_docx(canonical_assessment)
        doc = docx.Document(io.BytesIO(docx_bytes))
        full_text = _extract_all_doc_text(doc)

        assert "Summary judgement" in full_text
        # Check key phrases from summary judgement
        assert "On arithmetic this is the cleanest document" in full_text
        assert "The faults are in closure and navigation" in full_text

        # BOTTOM LINE callout
        assert "BOTTOM LINE" in full_text
        assert canonical_assessment.bottom_line in full_text
        assert "Fix conclusion 3 on printed page 28 before reissue" in full_text

    def test_the_four_baseline_measures_table(
        self,
        canonical_assessment: AssessmentData,
    ) -> None:
        """Section 3: The 4 baseline measures table has 4 rows with PASS/FAIL and reasons."""
        docx_bytes = generate_ale_review_docx(canonical_assessment)
        doc = docx.Document(io.BytesIO(docx_bytes))
        full_text = _extract_all_doc_text(doc)

        assert "The four baseline measures" in full_text
        assert "States the purpose of the report explicitly" in full_text
        assert "Procedure detailed enough for another competent party" in full_text
        assert "Conclusions are conclusions, not results and not discussion" in full_text
        assert "Recommendations name an owner and a date" in full_text

        # Baseline score
        assert "Baseline score:" in full_text
        assert "1/4 (1 of 4 passing)" in full_text

        # Find baseline table
        baseline_table = None
        for t in doc.tables:
            first_row_text = " ".join(c.text for c in t.rows[0].cells)
            if "Measure" in first_row_text and "Result" in first_row_text:
                baseline_table = t
                break

        assert baseline_table is not None
        assert len(baseline_table.rows) == 5  # 1 header + 4 baseline rows
        # Check pass/fail cells
        cell_results = [baseline_table.rows[i].cells[1].text.strip() for i in range(1, 5)]
        assert cell_results == ["FAIL", "PASS", "FAIL", "FAIL"]

    def test_blockers_callout_section(
        self,
        canonical_assessment: AssessmentData,
    ) -> None:
        """Section 4: Blockers are rendered in formatted callouts with all 5 required fields."""
        docx_bytes = generate_ale_review_docx(canonical_assessment)
        doc = docx.Document(io.BytesIO(docx_bytes))
        full_text = _extract_all_doc_text(doc)

        assert "Blockers" in full_text
        assert "Blocker 1:" in full_text
        assert "Blocker 2:" in full_text
        assert "THE CONCLUSIONS GIVE THE CRITICALITY 2 BAND" in full_text
        assert "THE DOCUMENT DOES NOT RECORD ITS OWN REVISION" in full_text

        # Check required blocker fields
        assert "Where:" in full_text
        assert "§7.1 conclusion 3, printed page 28." in full_text
        assert "What it says:" in full_text
        assert "What the body has:" in full_text
        assert "Why it matters:" in full_text
        assert "What would fix it:" in full_text

    def test_scorecard_summary_precedes_blockers(
        self,
        canonical_assessment: AssessmentData,
    ) -> None:
        docx_bytes = generate_ale_review_docx(canonical_assessment)
        full_text = _extract_all_doc_text(docx.Document(io.BytesIO(docx_bytes)))
        assert _section_position(full_text, "Scorecard") < _section_position(full_text, "Blockers")

    def test_repeated_control_findings_are_summarized(self) -> None:
        class Issue:
            included_in_report = True
            category = "LAYOUT"
            type = "UNCONTROLLED_PAGE"
            severity = "HIGH"
            message = "Footer control is missing."
            page_number = 1
            evidence = {"suggested_fix": "Apply one controlled footer template."}

        result = assessment_from_document_findings(object(), [Issue(), Issue()])
        assert len(result.blockers) == 1
        assert "2 uncontrolled page" in result.blockers[0].what_it_says

    def test_mixed_layout_control_findings_share_one_summary(self) -> None:
        class Issue:
            included_in_report = True
            category = "LAYOUT"
            severity = "MAJOR"
            page_number = 1
            message = "Document-control field is absent."
            evidence = {}

            def __init__(self, issue_type: str) -> None:
                self.type = issue_type

        issues = [Issue("UNCONTROLLED_PAGE") for _ in range(20)]
        issues.extend(Issue("MISSING_DOCUMENT_NUMBER") for _ in range(49))
        result = assessment_from_document_findings(object(), issues)
        assert len(result.blockers) == 1
        assert "69 related finding(s)" in result.blockers[0].what_it_says
        assert "20 uncontrolled page" in result.blockers[0].what_it_says
        assert "49 missing document number" in result.blockers[0].what_it_says

    def test_stale_standard_missing_bibliography_is_excluded(self) -> None:
        class Issue:
            included_in_report = True
            category = "STANDARD_TRACEABILITY"
            type = "STANDARD_NOT_IN_BIBLIOGRAPHY"
            severity = "HIGH"
            message = "ISO 9001 is missing."
            page_number = 1
            evidence = {}

        result = assessment_from_document_findings(object(), [Issue()])
        assert not result.blockers
        assert not result.major_findings

    def test_should_fix_in_next_revision_table(
        self,
        canonical_assessment: AssessmentData,
    ) -> None:
        """Section 5: Major findings table contains all 12 canonical findings."""
        docx_bytes = generate_ale_review_docx(canonical_assessment)
        doc = docx.Document(io.BytesIO(docx_bytes))
        full_text = _extract_all_doc_text(doc)

        assert "Should fix in the next revision" in full_text

        # Find major findings table
        maj_table = None
        for t in doc.tables:
            hdr_text = " ".join(c.text for c in t.rows[0].cells)
            if "No." in hdr_text and "Finding" in hdr_text and "What would fix it" in hdr_text:
                maj_table = t
                break

        assert maj_table is not None
        # 1 header row + 12 finding rows = 13 rows
        assert len(maj_table.rows) == 13
        assert len(canonical_assessment.major_findings) == 12

        # Check specific findings in table
        assert "front matter no longer navigates to the body" in full_text
        assert "Twenty-one of the fifty-six pages in the file carry no page number" in full_text
        assert "There is no reference section anywhere in the document" in full_text
        assert "The purpose of the report is never stated" in full_text
        assert "No statement of what the report contains and in what order" in full_text
        assert "None of the ten recommendations at §7.2 names an owner" in full_text
        assert "The Criticality 1 band is defined two ways" in full_text
        assert "Appendix C reports corrosion rates to nine decimal places" in full_text
        assert "§4 items 6 and 7 say the same thing" in full_text

    def test_language_and_mechanics_table(
        self,
        canonical_assessment: AssessmentData,
    ) -> None:
        """Section 6: Language and mechanics table contains all 15 language findings by page."""
        docx_bytes = generate_ale_review_docx(canonical_assessment)
        doc = docx.Document(io.BytesIO(docx_bytes))
        full_text = _extract_all_doc_text(doc)

        assert "Language and mechanics, by page" in full_text

        lang_table = None
        for t in doc.tables:
            hdr_text = " ".join(c.text for c in t.rows[0].cells)
            if "Page" in hdr_text and "As Written" in hdr_text and "Suggested" in hdr_text:
                lang_table = t
                break

        assert lang_table is not None
        # 1 header row + 15 language rows = 16 rows
        assert len(lang_table.rows) == 16
        assert len(canonical_assessment.language_findings) == 15

        # Check sample entries
        assert "202 7" in full_text
        assert "nine-teen" in full_text
        assert "comprising by" in full_text

    def test_demonstration_rewrite_section(
        self,
        canonical_assessment: AssessmentData,
    ) -> None:
        """Section 7: Demonstration rewrite compares original text against 8 single sentences."""
        docx_bytes = generate_ale_review_docx(canonical_assessment)
        doc = docx.Document(io.BytesIO(docx_bytes))
        full_text = _extract_all_doc_text(doc)

        assert "Demonstration rewrite" in full_text
        assert "§7.1 Conclusions" in full_text
        assert "AS WRITTEN (PRINTED PAGE 28, ABRIDGED)" in full_text
        assert "Faults against Budinski rules:" in full_text
        assert "DEMONSTRATION — THE SAME CONTENT AS CONCLUSIONS" in full_text

        # Check demonstration sentences
        assert (
            "No component of the 175 assessed has exhausted its remaining wall thickness"
            in full_text
        )
        assert "158 components, 90 per cent of those assessed" in full_text
        assert "One component, 35-V-101 SHELL, is classified as Criticality 2" in full_text
        assert "Eight numbered single sentences. No table reference" in full_text

    def test_scorecard_section_all_41_items(
        self,
        canonical_assessment: AssessmentData,
    ) -> None:
        """Section 8: Scorecard has group summary and all 41 items across 4 groups."""
        docx_bytes = generate_ale_review_docx(canonical_assessment)
        doc = docx.Document(io.BytesIO(docx_bytes))
        full_text = _extract_all_doc_text(doc)

        assert "Scorecard" in full_text
        assert "Group I: Technical Content" in full_text
        assert "Group II: Style" in full_text
        assert "Group III: Report Mechanics" in full_text
        assert "Group IV: Conclusions & Craft" in full_text
        assert "Overall Average" in full_text

        # Check 41 item codes in document text
        for idx in range(1, 10):
            assert f"I.{idx}" in full_text
        for idx in range(1, 12):
            assert f"II.{idx}" in full_text
        for idx in range(1, 12):
            assert f"III.{idx}" in full_text
        for idx in range(1, 11):
            assert f"IV.{idx}" in full_text

        # Verify rework items indicator (score <= 2)
        assert "(Rework)" in full_text

    def test_what_this_document_does_well(
        self,
        canonical_assessment: AssessmentData,
    ) -> None:
        """Section 9: 'What this document does well' lists all 8 positive observations."""
        docx_bytes = generate_ale_review_docx(canonical_assessment)
        doc = docx.Document(io.BytesIO(docx_bytes))
        full_text = _extract_all_doc_text(doc)

        assert "What this document does well" in full_text
        assert "Every count in the document reconciles." in full_text
        assert "The priority counts can be re-derived independently" in full_text
        assert "Table 3-1 and Table 3-2 together are the best single artefact" in full_text
        assert "§3.1 states the tie-breaking rules explicitly" in full_text
        assert "§4 gives fifteen numbered assumptions" in full_text
        assert "The recommendations carry dates." in full_text

    def test_limits_of_this_review_and_reviewscore(
        self,
        canonical_assessment: AssessmentData,
    ) -> None:
        """Section 10: 'Limits of this review' and REVIEWSCORE string are present."""
        docx_bytes = generate_ale_review_docx(canonical_assessment)
        doc = docx.Document(io.BytesIO(docx_bytes))
        full_text = _extract_all_doc_text(doc)

        assert "Limits of this review" in full_text
        assert "This is a technical writing review." in full_text
        assert "No figure in this report is asserted to be incorrect." in full_text
        assert "Anonymise before it becomes training material." in full_text

        # REVIEWSCORE summary string
        assert canonical_assessment.review_score_string in full_text
        assert "REVIEWSCORE | doc=" in full_text

    def test_docx_xml_styling_attributes(
        self,
        canonical_assessment: AssessmentData,
    ) -> None:
        """Verify Word XML attributes: tblHeader, cantSplit, and cell shading."""
        docx_bytes = generate_ale_review_docx(canonical_assessment)
        doc = docx.Document(io.BytesIO(docx_bytes))

        # Check for tblHeader on multi-row tables
        found_tbl_header = False
        found_cant_split = False

        for table in doc.tables:
            for row in table.rows:
                tr_xml = row._tr.xml
                if "w:tblHeader" in tr_xml:
                    found_tbl_header = True
                if "w:cantSplit" in tr_xml:
                    found_cant_split = True

        assert found_tbl_header, "Expected at least one table to have w:tblHeader repeating header"
        assert found_cant_split, "Expected table rows to have w:cantSplit row protection"

    def test_executive_report_sections_and_no_raw_alert_dump(
        self,
        canonical_assessment: AssessmentData,
    ) -> None:
        doc = docx.Document(io.BytesIO(generate_ale_review_docx(canonical_assessment)))
        full_text = _extract_all_doc_text(doc)
        for section in (
            "DOCUMENT REVIEW ENGINEERING",
            "Summary judgement",
            "BOTTOM LINE",
            "The four baseline measures",
            "Blockers",
            "Should fix in the next revision",
            "Language and mechanics, by page",
            "Demonstration rewrite",
            "Scorecard",
            "What this document does well",
            "Limits of this review",
        ):
            assert section in full_text
        assert '"evidence":' not in full_text
        assert "IssueRead(" not in full_text

    def test_styling_utilities_emit_valid_xml(self) -> None:
        doc = docx.Document()
        cell = create_callout_box(doc, "BLOCKER", ["Correct the source evidence."])
        set_cell_shading(cell, "#F8FAFC")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "measure"
        table.cell(0, 1).text = "result"
        format_table_header(table.rows[0], [2.0, 1.0])
        output = io.BytesIO()
        doc.save(output)
        parsed = docx.Document(io.BytesIO(output.getvalue()))
        xml = " ".join(table._tbl.xml for table in parsed.tables)
        assert "w:shd" in xml
        assert "w:tcBorders" in xml
        assert "w:tblHeader" in xml

    def test_report_synthesizer_is_deterministic(self) -> None:
        synthesizer = ReportSynthesizer()
        metadata = {"document_reviewed": "Inspection_Report.pdf"}
        scorecard = {"baseline_score": 2}
        findings = [{"type": "TOC_DRIFT", "message": "Contents page drift"}]
        first = synthesizer.generate_summary_judgement(findings, scorecard, metadata)
        second = synthesizer.generate_summary_judgement(findings, scorecard, metadata)
        assert first == second
        assert "Inspection_Report.pdf" in first
        assert "DOC-001" not in first
        assert "Rev A" not in first

    def test_export_adapter_accepts_persisted_computed_scorecard_fields(
        self, canonical_assessment: AssessmentData
    ) -> None:
        persisted = canonical_assessment.scorecard.model_dump(mode="json")
        adapted = assessment_from_document_findings(
            type("Document", (), {"original_filename": "fixture.pdf", "id": "doc-1"})(),
            [],
            scorecard_data=persisted,
        )
        assert adapted.scorecard.baseline_score == canonical_assessment.scorecard.baseline_score

    def test_edge_cases_empty_findings(self) -> None:
        """Report generates cleanly when blockers or findings lists are empty."""
        score_item_pass = ScoreItem(name="Test", score=5, note="Good")
        tech = TechnicalContentGroup(
            message_clear=score_item_pass,
            logical_approach=score_item_pass,
            adequate_research=score_item_pass,
            adequate_comparison=score_item_pass,
            conclusions_supported=score_item_pass,
            value_stated=score_item_pass,
            objective_met=score_item_pass,
            original_free_of_plagiarism=score_item_pass,
            timely=score_item_pass,
        )
        style = StyleGroup(
            objective_tone=score_item_pass,
            sections_logical=score_item_pass,
            readership_level=score_item_pass,
            free_of_jargon=score_item_pass,
            english_usage=score_item_pass,
            concise=score_item_pass,
            interesting=score_item_pass,
            free_of_personal_opinion=score_item_pass,
            no_over_explain=score_item_pass,
            standard_writing_practice=score_item_pass,
            layout_and_whitespace=score_item_pass,
        )
        mech = ReportMechanicsGroup(
            sufficient_background=score_item_pass,
            purpose_of_work_clear=score_item_pass,
            objective_of_work_clear=score_item_pass,
            purpose_of_report_clear=score_item_pass,
            objective_of_report_clear=score_item_pass,
            format_stated=score_item_pass,
            work_referenced=score_item_pass,
            experimental_steps_outlined=score_item_pass,
            adequate_detail_to_repeat=score_item_pass,
            free_of_trade_names=score_item_pass,
            test_standards_cited=score_item_pass,
        )
        cc = ConclusionsAndCraftGroup(
            results_clearly_stated=score_item_pass,
            results_free_of_discussion=score_item_pass,
            graphs_and_tables_proper=score_item_pass,
            sufficient_results=score_item_pass,
            discussion_relates_to_others=score_item_pass,
            discussion_length_appropriate=score_item_pass,
            conclusions_follow_from_results=score_item_pass,
            conclusions_clear=score_item_pass,
            references_properly_attributed=score_item_pass,
            sentence_paragraph_length=score_item_pass,
        )
        base = BaselineMeasures(
            purpose_distinct_from_objective=True,
            procedure_repeatable=True,
            conclusions_valid=True,
            recommendations_actionable=True,
            reasons={
                "purpose_distinct_from_objective": "Clearly stated.",
                "procedure_repeatable": "All formulas provided.",
                "conclusions_valid": "Valid numbered sentences.",
                "recommendations_actionable": "Owner and date given.",
            },
        )
        scorecard = BudinskiScorecard(
            technical_content=tech,
            style=style,
            report_mechanics=mech,
            conclusions_and_craft=cc,
            group_i_average=5.0,
            group_ii_average=5.0,
            group_iii_average=5.0,
            group_iv_average=5.0,
            overall_average=5.0,
            baseline_measures=base,
            blockers_count=0,
            majors_count=0,
            minors_count=0,
            review_score_string="REVIEWSCORE | perfect pass",
        )

        data = AssessmentData(
            title="Test Document Review",
            subtitle="Automated Verification Report",
            header_title="DOCUMENT REVIEW · QA",
            running_header="Test Document Rev 0",
            metadata=AssessmentMetadata(
                document_reviewed="DOC-TEST-001 Rev 0",
                type_of_review="Technical writing review",
                basis="Appendix 12",
                scoring="1-5 scale",
                note="Automated unit test run",
                not_covered="Engineering calculations",
            ),
            summary_judgement=["Clean test document."],
            bottom_line="Document passes all writing quality gates.",
            baseline_measures=base,
            blockers=[],
            major_findings=[],
            language_findings=[],
            demonstration_rewrite=None,
            scorecard=scorecard,
            what_it_does_well=["Excellent structure throughout."],
            limits_of_review=["Writing review only."],
            review_score_string="REVIEWSCORE | perfect pass",
        )

        docx_bytes = generate_ale_review_docx(data)
        doc = docx.Document(io.BytesIO(docx_bytes))
        full_text = _extract_all_doc_text(doc)

        assert "No blocking findings identified." in full_text
        assert "No major findings identified." in full_text
        assert "No language or mechanics findings identified." in full_text
        assert "Baseline score:" in full_text
        assert "4/4 (4 of 4 passing)" in full_text


class TestTemplateGenerator:
    """Verification suite for create_review_report_template."""

    def test_create_review_report_template_bytes(self) -> None:
        """Template generator returns non-empty docx binary stream."""
        template_bytes = create_review_report_template()
        assert isinstance(template_bytes, bytes)
        assert len(template_bytes) > 2000
        assert template_bytes.startswith(b"PK\x03\x04")

        doc = docx.Document(io.BytesIO(template_bytes))
        assert len(doc.paragraphs) > 0
        full_text = " ".join(p.text for p in doc.paragraphs)
        assert "DOCUMENT REVIEW TEMPLATE" in full_text

    def test_create_review_report_template_file_save(self, tmp_path: Path) -> None:
        """Template generator writes file to disk when output_path is provided."""
        target_path = tmp_path / "test_template.docx"
        template_bytes = create_review_report_template(target_path)
        assert target_path.exists()
        assert target_path.stat().st_size == len(template_bytes)


class TestBilingualStrictIsolation:
    """Verification suite for strict language purity (Anti-Bahasa Belang)."""

    def test_strict_english_purity(self, canonical_assessment: AssessmentData) -> None:
        """Language=en must have full English boilerplate and zero Indonesian words."""
        docx_bytes = generate_ale_review_docx(canonical_assessment, language=ReportLanguage.ENGLISH)
        doc = docx.Document(io.BytesIO(docx_bytes))
        full_text = _extract_all_doc_text(doc)

        # Expected English structural anchors
        assert "DOCUMENT REVIEW ENGINEERING" in full_text
        assert "Document reviewed" in full_text
        assert "The four baseline measures" in full_text
        assert "Scorecard" in full_text
        assert "Blockers" in full_text
        assert "Should fix in the next revision" in full_text
        assert "Language and mechanics, by page" in full_text
        assert "What this document does well" in full_text
        assert "Limits of this review" in full_text
        assert "Group I: Technical Content" in full_text
        assert "Group II: Style" in full_text
        assert "Group III: Report Mechanics" in full_text
        assert "Group IV: Conclusions & Craft" in full_text
        assert "Overall Average" in full_text
        assert "PASS" in full_text

        # Strict absence of Indonesian boilerplate keywords
        indonesian_forbidden = [
            "Dokumen yang ditinjau",
            "Empat ukuran dasar",
            "Kartu nilai",
            "Penghalang",
            "Perbaiki pada revisi berikutnya",
            "Bahasa dan mekanika, menurut halaman",
            "Keunggulan dokumen ini",
            "Batasan tinjauan ini",
            "Grup I: Konten Teknis",
            "Rata-rata Keseluruhan",
            "DITOLAK",
            "LULUS",
            "Catatan Peninjau",
            "Parameter Daftar Periksa",
        ]
        for word in indonesian_forbidden:
            assert word not in full_text, (
                f"Found unexpected Indonesian string in English report: '{word}'"
            )

    def test_strict_indonesian_purity(self, canonical_assessment: AssessmentData) -> None:
        """Language=id must have full Indonesian boilerplate and zero English boilerplate."""
        canonical_assessment.metadata.not_covered = "Kecukupan rekayasa sendiri."
        docx_bytes = generate_ale_review_docx(
            canonical_assessment, language=ReportLanguage.INDONESIAN
        )
        doc = docx.Document(io.BytesIO(docx_bytes))
        full_text = _extract_all_doc_text(doc)

        # Expected Indonesian structural anchors
        assert "REKAYASA TINJAUAN DOKUMEN" in full_text
        assert "Dokumen yang ditinjau" in full_text
        assert "Empat ukuran dasar" in full_text
        assert "Kartu nilai" in full_text
        assert "Penghalang" in full_text
        assert "Perbaiki pada revisi berikutnya" in full_text
        assert "Bahasa dan mekanika, menurut halaman" in full_text
        assert "Keunggulan dokumen ini" in full_text
        assert "Batasan tinjauan ini" in full_text
        assert "Grup I: Konten Teknis" in full_text
        assert "Grup II: Gaya" in full_text
        assert "Grup III: Mekanika Laporan" in full_text
        assert "Grup IV: Kesimpulan & Keterampilan" in full_text
        assert "Rata-rata Keseluruhan" in full_text
        assert "LULUS" in full_text

        # Strict absence of English boilerplate phrases
        english_forbidden = [
            "Document reviewed",
            "The four baseline measures",
            "Should fix in the next revision",
            "Language and mechanics, by page",
            "What this document does well",
            "Limits of this review",
            "Group I: Technical Content",
            "Group II: Style",
            "Overall Average",
            "Items requiring rework",
            "Checklist Parameter",
            "Reviewer Note",
        ]
        for phrase in english_forbidden:
            assert phrase not in full_text, (
                f"Found untranslated English phrase in Indonesian report: '{phrase}'"
            )


class TestLibreOfficePDFConversion:
    """Verification suite for headless LibreOffice DOCX to PDF conversion."""

    def test_convert_docx_to_pdf_missing_file(self, tmp_path: Path) -> None:
        """Attempting to convert non-existent docx raises FileNotFoundError."""
        with pytest.raises(FileNotFoundError):
            convert_docx_to_pdf(tmp_path / "nonexistent.docx", tmp_path)

    def test_convert_docx_to_pdf_mock_subprocess(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """convert_docx_to_pdf invokes headless command and verifies PDF output existence."""
        import subprocess

        dummy_docx = tmp_path / "test.docx"
        dummy_docx.write_bytes(b"dummy docx content")
        out_pdf = tmp_path / "test.pdf"

        def mock_run(cmd: list[str], **kwargs: Any) -> subprocess.CompletedProcess[str]:
            assert "--headless" in cmd
            assert "--convert-to" in cmd
            assert "pdf" in cmd
            assert "--outdir" in cmd
            assert str(tmp_path) in cmd
            assert str(dummy_docx) in cmd
            out_pdf.write_bytes(b"%PDF-1.4 mock pdf content")
            return subprocess.CompletedProcess(cmd, returncode=0, stdout="converted", stderr="")

        monkeypatch.setattr(subprocess, "run", mock_run)
        result_path = convert_docx_to_pdf(dummy_docx, tmp_path)
        assert result_path == out_pdf
        assert result_path.exists()
        assert result_path.read_bytes().startswith(b"%PDF-")

    def test_convert_docx_to_pdf_error_handling(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        """convert_docx_to_pdf raises LibreOfficeConversionError on non-zero exit code."""
        import subprocess

        dummy_docx = tmp_path / "test_err.docx"
        dummy_docx.write_bytes(b"dummy")

        def mock_run_err(cmd: list[str], **kwargs: Any) -> subprocess.CompletedProcess[str]:
            return subprocess.CompletedProcess(cmd, returncode=1, stdout="", stderr="crash")

        monkeypatch.setattr(subprocess, "run", mock_run_err)
        with pytest.raises(LibreOfficeConversionError):
            convert_docx_to_pdf(dummy_docx, tmp_path)
