"""Unit and integration tests for Report Export Service (Annotated PDF and DOCX Review Report)."""

from __future__ import annotations

import io
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from uuid import uuid4

import docx
import pypdf
import pytest
from httpx import ASGITransport, AsyncClient

from domain.enums import DocumentStatus, IssueCategory, Severity
from main import app
from models.document import Document
from models.issue import Issue
from services.export import export_annotated_pdf
from services.report import build_review_report
from services.storage.local import LocalStorage


def _make_sample_models(
    tmp_path: Path | str,
) -> tuple[Document, list[Issue]]:
    """Create sample Document and Issue models for export testing."""
    doc_id = uuid4()
    now = datetime.now(UTC)

    doc = Document(
        id=doc_id,
        original_filename="WELD-SPEC-001.pdf",
        safe_filename="weld_spec_001",
        media_type="application/pdf",
        size_bytes=1024,
        sha256="e3b0c44298fc1c149afbf4c8996fb92427ae41e4649b934ca495991b7852b855",
        status=DocumentStatus.COMPLETED,
        storage_uri=f"file://documents/{doc_id}.pdf",
        canonical_pdf_uri=f"file://documents/{doc_id}.pdf",
        created_at=now,
        updated_at=now,
    )

    issue_trace = Issue(
        id=uuid4(),
        document_id=doc_id,
        category=IssueCategory.TRACEABILITY,
        type="TABLE_MATH_MISMATCH",
        severity=Severity.CRITICAL,
        confidence=0.98,
        message="Column sum mismatch in Table 2.1: stated 1,250.00 != computed 1,350.00",
        evidence={
            "kind": "TABLE_MATH",
            "stated_value": "1,250.00",
            "computed_value": "1,350.00",
            "delta": "100.00",
            "tolerance": "0.01",
            "suggestion": "Update stated sum in Table 2.1 to 1,350.00",
            "total_location": {
                "page_index": 0,
                "x0": 100.0,
                "y0": 200.0,
                "x1": 300.0,
                "y1": 250.0,
                "page_width": 612.0,
                "page_height": 792.0,
            },
        },
        included_in_report=True,
        reviewer_note="Confirmed with Lead Welding Engineer.",
        created_at=now,
        updated_at=now,
    )

    issue_ling = Issue(
        id=uuid4(),
        document_id=doc_id,
        category=IssueCategory.LINGUISTIC,
        type="SPELLING_ERROR",
        severity=Severity.LOW,
        confidence=0.92,
        message="Possible spelling error 'temprature'. Suggested replacement: 'temperature'.",
        evidence={
            "kind": "LINGUISTIC",
            "original_text": "temprature",
            "suggestion": "temperature",
            "location": {
                "page_index": 0,
                "x0": 50.0,
                "y0": 80.0,
                "x1": 150.0,
                "y1": 100.0,
                "page_width": 612.0,
                "page_height": 792.0,
            },
        },
        included_in_report=True,
        reviewer_note=None,
        created_at=now,
        updated_at=now,
    )

    issue_excluded = Issue(
        id=uuid4(),
        document_id=doc_id,
        category=IssueCategory.LINGUISTIC,
        type="GRAMMAR_STYLE",
        severity=Severity.INFO,
        confidence=0.80,
        message="Passive voice usage detected.",
        evidence={"suggestion": "Use active voice."},
        included_in_report=False,
        reviewer_note="Excluded: acceptable passive voice in engineering specification.",
        created_at=now,
        updated_at=now,
    )

    return doc, [issue_trace, issue_ling, issue_excluded]


def test_export_annotated_pdf(tmp_path: Path) -> None:
    """Export annotated PDF embeds bounding boxes and callout popups without error."""
    storage = LocalStorage(root=tmp_path)
    doc, issues = _make_sample_models(tmp_path)
    included = [i for i in issues if i.included_in_report]

    pdf_bytes = export_annotated_pdf(doc, included, storage)
    assert len(pdf_bytes) > 0
    assert pdf_bytes.startswith(b"%PDF-")

    # Validate PDF structure with pypdf
    reader = pypdf.PdfReader(io.BytesIO(pdf_bytes))
    assert len(reader.pages) >= 1
    page = reader.pages[0]
    assert "/Annots" in page


def test_build_review_report_docx(tmp_path: Path) -> None:
    """DOCX review report builder creates structured Word document."""
    doc, issues = _make_sample_models(tmp_path)

    docx_bytes = build_review_report(doc, issues)
    assert len(docx_bytes) > 0

    word_doc = docx.Document(io.BytesIO(docx_bytes))
    full_text = "\n".join(p.text for p in word_doc.paragraphs)
    all_table_text = "\n".join(
        " ".join(cell.text for cell in row.cells)
        for table in word_doc.tables
        for row in table.rows
    )

    # Verify key sections
    assert "DOCUMENT REVIEW ENGINEERING" in full_text
    assert doc.original_filename in full_text
    assert "Summary judgement" in full_text
    assert "Scorecard" in full_text
    assert "Blockers" in full_text
    assert "Next revision findings" in full_text
    assert "The four baseline measures" in full_text
    assert "Scope and limitations" in full_text

    # Verify blocker content and reviewer note
    assert "TABLE_MATH_MISMATCH" in all_table_text
    assert "Reviewer note: Confirmed with Lead Welding Engineer." in all_table_text
    assert "Update stated sum in Table 2.1 to 1,350.00" in all_table_text

    # Verify excluded finding is NOT in report
    assert "Passive voice usage detected." not in full_text

    # Verify scorecard table
    assert len(word_doc.tables) >= 1
    table = word_doc.tables[0]
    cell_texts = [cell.text for row in table.rows for cell in row.cells]
    assert "Included findings" in cell_texts
    assert "2" in cell_texts  # 2 included findings
    assert "Blockers" in cell_texts
    assert "1" in cell_texts  # 1 critical blocker


@pytest.mark.anyio
async def test_export_endpoints(tmp_path: Path) -> None:
    """HTTP export endpoint handles PDF and DOCX formats and rejects legacy formats."""
    from core.dependencies import get_storage
    from db.session import get_session

    doc, issues = _make_sample_models(tmp_path)
    storage = LocalStorage(root=tmp_path)

    class MockAsyncSession:
        async def execute(self, stmt: Any) -> Any:
            class MockResult:
                def scalar_one_or_none(self) -> Any:
                    return doc

                def scalars(self) -> Any:
                    class MockScalars:
                        def all(self) -> list[Any]:
                            return issues

                    return MockScalars()

            return MockResult()

    async def override_get_session() -> Any:
        yield MockAsyncSession()

    app.dependency_overrides[get_session] = override_get_session
    app.dependency_overrides[get_storage] = lambda: storage

    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
            # 1. PDF
            res_pdf = await client.get(f"/api/v1/documents/{doc.id}/export?format=pdf")
            assert res_pdf.status_code == 200
            assert res_pdf.headers["content-type"] == "application/pdf"
            assert "attachment" in res_pdf.headers["content-disposition"]
            assert "weld_spec_001_annotated.pdf" in res_pdf.headers["content-disposition"]
            assert res_pdf.content.startswith(b"%PDF-")

            # 2. DOCX
            res_docx = await client.get(f"/api/v1/documents/{doc.id}/export?format=docx")
            assert res_docx.status_code == 200
            assert (
                res_docx.headers["content-type"]
                == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            )
            assert "attachment" in res_docx.headers["content-disposition"]
            assert "weld_spec_001_review.docx" in res_docx.headers["content-disposition"]

            # 3. Legacy formats (xlsx, csv, json) rejected with 422
            res_xlsx = await client.get(f"/api/v1/documents/{doc.id}/export?format=xlsx")
            assert res_xlsx.status_code == 422

            res_csv = await client.get(f"/api/v1/documents/{doc.id}/export?format=csv")
            assert res_csv.status_code == 422

            res_json = await client.get(f"/api/v1/documents/{doc.id}/export?format=json")
            assert res_json.status_code == 422
    finally:
        app.dependency_overrides.clear()
